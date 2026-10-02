import streamlit as st
from datetime import datetime
import os
import openpyxl
import smtplib
from email.message import EmailMessage

# Configurazione pagina Streamlit
st.set_page_config(
    page_title="ATS LIGURIA - Sistema di Gestione Qualità Aziendale",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# --- CONFIGURAZIONE FILE EXCEL ---
NOME_FILE_EXCEL = "2026_09_02_Prova_CL_Antincendio.xlsx"
NOME_FOGLIO_TARGET = "REGISTRAZIONE"

# --- OPZIONI E MENU A TENDINA ---
OPZIONI_DIPARTIMENTI = [
    "EMERGENZA E ACCETTAZIONE",
    "CHIRURGICO",
    "ORTOPEDICO TRAUMATOLOGICO",
    "MEDICO",
    "NEUROSCIENZE",
    "ONCOLOGICO",
    "RADIOLOGIA DEI SERVIZI",
    "PROGRAMMAZIONE SANITARIA E GOVERNO CLINICO",
    "PREVENZIONE",
    "ATTIVITÀ TERRITORIALI DELLA RIABILITAZIONE",
    "SALUTE MENTALE DELLE DIPENDENZE",
    "TECNICO AMMINISTRATIVO",
]

OPZIONI_SC_SSD = [
    "A.D.I. INFERMIERI SAVONA",
    "DISTRETTO FINALESE",
    "S.C . FARMACEUTICA",
    "S.C. 118 EMERGENZA TERRITORIALE",
    "S.C. ANATOMIA PATOLOGICA",
    "S.C. ANESTESIA E RIANIMAZIONE LEVANTE",
    "S.C. ANESTESIA E RIANIMAZIONE PONENTE",
    "S.C. ANGIOGRAFIA INTERVENTISTICA",
    "S.C. CARDIOLOGIA E UTIC PONENTE",
    "S.C. CHIRURGIA DELLA MANO",
    "S.C. CHIRURGIA GENERALE INDIRIZZO ENDOCRINO METABOLICO",
    "S.C. CHIRURGIA PLASTICA",
    "S.C. CHIRURGIA PROTESICA",
    "S.C. DIREZIONE MEDICA",
    "S.C. DIREZIONE PROFESSIONI SANITARIE",
    "S.C. DISTRETTO ALBENGANESE",
    "S.C. DISTRETTO DELLE BORMIDE",
    "S.C. DISTRETTO FINALESE",
    "S.C. DISTRETTO SAVONESE",
    "S.C. FORMAZIONE, INNOVAZIONE, SVILUPPO PROFESSIONALE",
    "S.C. GASTROENTEROLOGIA",
    "S.C. GINECOLOGIA E OSTETRICIA LEVANTE",
    "S.C. IGIENE E SANITÀ PUBBLICA",
    "S.C. LABORATORIO DI PATOLOGIA CLINICA",
    "S.C. MALATTIE INFETTIVE",
    "S.C. MEDICINA INTERNA 1 LEVANTE",
    "S.C. MEDICINA INTERNA 2 LEVANTE",
    "S.C. MEDICINA INTERNA PONENTE",
    "S.C. MEDICINA TRASFUSIONALE E IMMUNOEMATOLOGIA",
    "S.C. NEFROLOGIA E DIALISI",
    "S.C. NEUROCHIRURGIA",
    "S.C. NEUROLOGIA LEVANTE",
    "S.C. NEUROLOGIA PONENTE",
    "S.C. NEURORADIOLOGIA",
    "S.C. ORTOPEDIA E TRAUMATOLOGIA LEVANTE",
    "S.C. ORTOPEDIA E TRAUMATOLOGIA PONENTE",
    "S.C. PATRIMONIO E GESTIONE TECNICA",
    "S.C. PNEUMOLOGIA",
    "S.C. PRONTO SOCCORSO E MEDICINA D'URGENZA LEVANTE",
    "S.C. PRONTO SOCCORSO E MEDICINA D'URGENZA PONENTE",
    "S.C. PROVVEDITORATO",
    "S.C. PSICHIATRIA TERRITORIALE",
    "S.C. RADIOLOGIA LEVANTE",
    "S.C. RADIOTERAPIA",
    "S.C. RECUPERO E RIEDUCAZIONE FUNZIONALE",
    "S.C. SERVIZIO PSICHIATRICO CURE RIABILITATIVE - CENTRO DISTURBI ALIMENTAZIONE",
    "S.C. SERVIZIO PSICHIATRICO DIAGNOSI E CURA - S.P.D.C.",
    "S.C. UNITÀ SPINALE",
    "S.C. UROLOGIA",
    "S.S.D. CHIRURGIA ORTOPEDICA SETTICA",
    "S.S.D. CHIRURGIA VERTEBRALE",
    "S.S.D. ENDOSCOPIA DIGESTIVA SAVONA",
    "S.S.D. MALATTIE DEL SANGUE E NEOPLASIE DEGLI ORGANI EMOLINFOPOIETICI",
    "S.S.D. MEDICINA DEL LAVORO",
    "S.S.D. MICROBIOLOGIA",
    "S.S.D. PSICOLOGIA CLINICA",
]

OPZIONI_SEDI = [
    "SAN PAOLO SAVONA",
    "SANTA CORONA PIETRA LIGURE",
    "SANTA MARIA MISERICORDIA ALBENGA",
    "SAN GIUSEPPE CAIRO",
]

OPZIONI_DESC_SEDE = ["TERRITORIO", "OSPEDALE"]
ESITI_STANDARD = ["SI", "NO", "N.P."]
SI_NO = ["SI", "NO"]

OPZIONI_PROBLEMI_ESTINTORI = [
    "",
    "NON ACCESSIBILE",
    "MANCA SEGNALETICA",
    "NON ADEGUATAMENTE ANCORATO",
    "MANCA LO SPINOTTO",
    "MANCA IL SIGILLO",
    "MANCA IL CARTELLINO",
    "MANICHETTA ROTTA/USURATA",
    "CONO DIFFUSORE NON INTEGRO",
    "MANOMETRO SOTTO PRESSIONE",
    "NON ESEGUITA VERIFICA SEMESTRALE DA PARTE DI DITTA",
]

OPZIONI_PROBLEMI_IDRANTI = [
    "",
    "NON ACCESSIBILE",
    "MANCA SEGNALETICA",
    "SAFE CRASH CON SEGNI DI ROTTURA",
    "MANCA LA LANCIA",
    "MANCA LA MANICHETTA",
    "MANCA IL CARTELLINO",
    "CARTELLINO NON VISIBILE",
    "MANCA IL SIGILLO",
    "NON COLLEGATO A RETE IDRICA",
    "LANCIA E MANICHETTA SCOLLEGATI TRA LORO",
    "NON ESEGUITA VERIFICA SEMESTRALE DA PARTE DI DITTA",
]

OPZIONI_REI = [
    "",
    "Tocca a pavimento",
    "Selettore di chiusura",
    "Non sgancia dai magneti",
    "Non aggancia ai magneti",
    "Non torna in totale chiusura anta",
    "Maniglia semplice",
    "Maniglione antipanico rotta o mal conservata",
    "Difficoltà di apertura",
    "Non eseguita verifica semestrale da parte della ditta",
    "Manca cartellino di verifica",
]

OPZIONI_NON_REI = [
    "",
    "Tocca a pavimento",
    "Non rimane in chiusura",
    "Difficoltà in apertura/anta bloccata",
    "Non si chiude correttamente anta",
    "Maniglia semplice",
    "Maniglione antipanico rotto,mal conservato o di difficle apertura",
    "Vetro rotto o venato",
    "Non eseguita verifica semestrale da parte della Ditta",
    "Manca il cartellino di verifica",
]

OPZIONI_USCITE = [
    "",
    "Materiale sanitario/scatole depositato davanti all'uscita",
    "Attrezzatura/macchinari elettromedicali che ostruiscono il passaggio",
    "Arredi/mobili/carrelli sanitari che ostacolano il passaggio",
    "Porta bloccata o non apribile",
    "Letto paziente in prossimità dell'uscita",
    "Barella presente davanti all'uscita",
    "Carrozzina/Sedia presente davanto all'uscita",
    "Bombole di ossigeno o altro gas medicale",
    "Incombro dovuto a lavori o manutenzione",
    "Materiale temporaneamente accatastato",
    "Rifiuti imballaggi sacchi biancheria presenti",
    "Uscita ostruita completamente",
    "Uscita ostruita parzialmente",
    "Manca il cartello di segnalazione su porta",
    "I cartelli indicano la direzione sbagliata",
]

OPZIONI_CARTELLO_LUMINOSO = ["SI", "NO", "N.P.", "Cartello rotto", "Cartello non illuminato"]
OPZIONI_LED_CARICA = ["ROSSO", "SPENTO"]
OPZIONI_NUM_DEPOSITI = ["NO", "1", "2", "3", "4", "+di 4"]
OPZIONI_SCAFFALI = ["APERTE", "CHIUSE"]

# --- INTESTAZIONE GRAFICA ---
col_logo1, col_logo2, col_logo3 = st.columns([1, 2, 1])
with col_logo2:
    logo_path = "ATS.LIGURIA_LOGO_.png"
    if os.path.exists(logo_path):
        st.image(logo_path, use_container_width=True)

st.markdown("<h3 style='text-align: center; color: #003366;'>SERVIZIO DI PREVENZIONE E PROTEZIONE</h3>", unsafe_allow_html=True)
st.markdown("<p style='text-align: center; font-style: italic; color: #475569;'>Sistema di Gestione Qualità Aziendale</p>", unsafe_allow_html=True)
st.markdown("<div style='background-color: #003366; padding: 10px; border-radius: 5px; text-align: center; color: white; font-weight: bold; margin-bottom: 20px;'>CHECK LIST - VERIFICHE DISPOSITIVI ANTINCENDIO</div>", unsafe_allow_html=True)

# --- FORM DI COMPILAZIONE ---
with st.form("form_ispezione"):
    
    # 0. DATI COMPILATORE
    st.markdown("#### 👤 DATI COMPILATORE")
    matricola = st.text_input("Matricola Operatore:")
    
    st.divider()
    
    # 1. UBICAZIONE
    st.markdown("#### 📍 UBICAZIONE")
    col1, col2 = st.columns(2)
    with col1:
        dipartimento = st.selectbox("Dipartimento*:", [""] + OPZIONI_DIPARTIMENTI)
        sede = st.selectbox("Sede Aziendale:", [""] + OPZIONI_SEDI)
        comune = st.text_input("Comune:")
    with col2:
        sc_ssd = st.selectbox("SC / SSD*:", [""] + OPZIONI_SC_SSD)
        desc_sede = st.selectbox("Descrizione Sede:", OPZIONI_DESC_SEDE)
        via = st.text_input("Via:")
        
    col_p1, col_p2, col_p3 = st.columns(3)
    with col_p1:
        padiglione = st.text_input("Padiglione:")
    with col_p2:
        piano = st.text_input("Piano:")
    with col_p3:
        lato = st.text_input("Lato:")

    st.divider()

    # 2. VERIFICA ESTINTORI
    st.markdown("#### 🧯 VERIFICA ESTINTORI (SI / NO / N.P.)")
    est1 = st.selectbox("1. Visibili e facilmente accessibili:", ESITI_STANDARD, index=0)
    est2 = st.selectbox("2. Posizione indicata con segnaletica:", ESITI_STANDARD, index=0)
    est3 = st.selectbox("3. Ancorato a muro o supporto portatile:", ESITI_STANDARD, index=0)
    est4 = st.selectbox("4. Presente spinotto di sicurezza e sigillo:", ESITI_STANDARD, index=0)
    est5 = st.selectbox("5. Cartellino di manutenzione compilato:", ESITI_STANDARD, index=0)
    est6 = st.selectbox("6. Manichetta/cono CO2 integro e libero:", ESITI_STANDARD, index=0)
    est7 = st.selectbox("7. Indicatore di pressione (estintore a polvere) nel settore verde:", ESITI_STANDARD, index=0)
    
    col_e1, col_e2 = st.columns(2)
    with col_e1:
        estintore_num = st.text_input("Estintore n.:")
    with col_e2:
        estintore_problema = st.selectbox("Specificare problema:", OPZIONI_PROBLEMI_ESTINTORI)
        
    note_estintori = st.text_input("Note / Dettagli Estintori:")

    st.divider()

    # 3. ALTRI PRESIDI ED IMPIANTI
    st.markdown("#### 💧 ALTRI PRESIDI ED IMPIANTI")
    idranti = st.selectbox("Idranti / Naspo:", ESITI_STANDARD, index=0)
    
    col_i1, col_i2 = st.columns(2)
    with col_i1:
        idrante_num = st.text_input("Idrante n.:")
    with col_i2:
        idrante_problema = st.selectbox("Specificare problema Idrante:", OPZIONI_PROBLEMI_IDRANTI)
    note_idranti = st.text_input("Note Idranti:")

    coperta = st.selectbox("Coperta Antifiamma presente:", ESITI_STANDARD, index=0)
    coperta_segnalata = st.selectbox("Dove è presente, è correttamente segnalata:", ESITI_STANDARD, index=0)
    coperta_integra = st.selectbox("La custodia della coperta è integra:", ESITI_STANDARD, index=0)
    note_coperta = st.text_input("Note Coperta Antifiamma:")

    fumi = st.selectbox("Rilevatori di fumo a soffitto sono integri:", ESITI_STANDARD, index=0)
    note_fumi = st.text_input("Rilevatori di fumo se NO integri specificare il dettaglio:")

    centralina = st.selectbox("Centralina di segnalazione di fumo sul display sono segnalati errori o guasti:", ESITI_STANDARD, index=0)
    note_centralina = st.text_input("Se SI specificare Centralina:")

    pulsanti = st.selectbox("Pulsanti di allarme sono visibili, utilizzabili e segnalati:", ESITI_STANDARD, index=0)
    pulsanti_privi = st.text_input("Se NO nota per segnalare quanti pulsanti privi di cartello segnalazione:")

    st.divider()

    # 4. PORTE E PORTE NON REI
    st.markdown("#### 🚪 PORTE E PORTE NON REI")
    rei_apertura = st.selectbox("Porte REI si aprono e chiudono con facilità:", ESITI_STANDARD, index=0)
    rei_maniglioni = st.selectbox("I maniglioni antipanico dove presenti sono integri e funzionanti:", ESITI_STANDARD, index=0)
    
    col_r1, col_r2, col_r3 = st.columns(3)
    with col_r1:
        rei_num_porte = st.text_input("Numero delle porte REI:")
    with col_r2:
        rei_num_ante = st.text_input("Numero delle ante REI:")
    with col_r3:
        rei_prob = st.selectbox("Problema riscontrato REI:", OPZIONI_REI)

    non_rei_esito = st.selectbox("Porte non REI apertura e chiusura avvengono con facilità:", ESITI_STANDARD, index=0)
    non_rei_num = st.text_input("Numero delle porte non REI:")
    non_rei_prob = st.selectbox("Se NO segnalare il dettaglio non REI:", OPZIONI_NON_REI)
    non_rei_ante = st.text_input("Numero ante non REI:")
    non_rei_stato = st.text_input("Stato integrità della porta non REI:")

    st.divider()

    # 5. USCITE DI SICUREZZA, PLANIMETRIE E LAMPADE
    st.markdown("#### 🚪 USCITE DI SICUREZZA, PLANIMETRIE E LAMPADE")
    uscite_esito = st.selectbox("Uscite di sicurezza sono libere da ostacoli e ben segnalate:", ESITI_STANDARD, index=0)
    
    col_u1, col_u2 = st.columns(2)
    with col_u1:
        uscite_prob = st.selectbox("Se NO specificare (Dettaglio):", OPZIONI_USCITE)
    with col_u2:
        uscite_altro = st.text_input("Altro (Uscite):")

    col_c1, col_c2 = st.columns(2)
    with col_c1:
        cartello_lum = st.selectbox("Cartello luminoso uscite di sicurezza è integro:", OPZIONI_CARTELLO_LUMINOSO)
    with col_c2:
        cartello_altro = st.text_input("Altro (Cartello):")

    planimetrie_pres = st.selectbox("Planimetrie e percorsi di esodo sono presenti:", ESITI_STANDARD, index=0)
    planimetrie_agg = st.selectbox("Se planimetrie presenti sono corrette e aggiornate:", ESITI_STANDARD, index=0)
    
    lampade_integre = st.selectbox("Le lampade di emergenza sono integre:", ESITI_STANDARD, index=0)
    lampade_rotte_num = st.text_input("Se NO specificare dettaglio N° lampade rotte:")

    carica_led_pres = st.selectbox("Carica batterie (se presente LED):", ESITI_STANDARD, index=0)
    col_l1, col_l2 = st.columns(2)
    with col_l1:
        led_lampada_num = st.text_input("Lampada N° (LED):")
    with col_l2:
        led_stato = st.selectbox("LED Stato:", [""] + OPZIONI_LED_CARICA)

    st.divider()

    # 6. LIQUIDI INFIAMMABILI ED ARMADI
    st.markdown("#### 🧪 LIQUIDI INFIAMMABILI ED ARMADI DI SICUREZZA")
    liquidi_pres = st.selectbox("Sono presenti liquidi infiammabili:", ESITI_STANDARD, index=0)
    col_lq1, col_lq2 = st.columns(2)
    with col_lq1:
        liquidi_litri = st.text_input("N° Litri (Liquidi):")
    with col_lq2:
        liquidi_stanze = st.text_input("N° Stanze (Liquidi):")

    armadi_integri = st.selectbox("Armadi di sicurezza stoccaggio prodotti infiammabili sono integri:", ESITI_STANDARD, index=0)
    armadi_note_libera = st.text_input("Se NO specificare (Nota libera Armadi):")

    st.divider()

    # 7. LOCALI DEPOSITO
    st.markdown("#### 📦 LOCALI DEPOSITO")
    depositi_num = st.selectbox("Numero depositi presenti:", OPZIONI_NUM_DEPOSITI)
    depositi_fumi = st.selectbox("Rilevatori di fumo presenti e integri (Depositi):", SI_NO)
    depositi_rei = st.selectbox("Porte REI tagliafuoco integre e chiuse (Depositi):", SI_NO)
    depositi_estintore = st.selectbox("Estintore presente e accessibile (Depositi):", SI_NO)
    
    col_sc1, col_sc2 = st.columns(2)
    with col_sc1:
        scaffali_presenza = st.selectbox("Scaffalature Presenza:", SI_NO)
    with col_sc2:
        scaffali_tipo = st.selectbox("Scaffalature Tipo:", OPZIONI_SCAFFALI)

    st.divider()

    # 8. BOMBOLE DI GAS MEDICALI
    st.markdown("#### 💨 BOMBOLE DI GAS MEDICALI")
    bombole_pres = st.selectbox("Sono presenti bombole di gas medicali:", SI_NO)
    col_b1, col_b2 = st.columns(2)
    with col_b1:
        bombole_num = st.text_input("N° Bombole:")
    with col_b2:
        bombole_litri = st.text_input("Litri N° (Bombole):")

    bombole_fumi = st.selectbox("Rilevatore di fumo presente (Bombole):", SI_NO)
    bombole_areato = st.selectbox("Locale ben areato (Bombole):", SI_NO)
    bombole_ancorate = st.selectbox("Bombole correttamente ancorate:", SI_NO)

    st.markdown("<br>", unsafe_allow_html=True)
    
    # Pulsante di Invio/Salvataggio pulito
    submitted = st.form_submit_button("Salva e invia")

    if submitted:
        if not dipartimento or not sc_ssd:
            st.error("⚠️ I campi 'Dipartimento' e 'SC/SSD' sono obbligatori!")
        else:
            try:
                if not os.path.exists(NOME_FILE_EXCEL):
                    wb_new = openpyxl.Workbook()
                    ws_new = wb_new.active
                    ws_new.title = NOME_FOGLIO_TARGET
                    wb_new.save(NOME_FILE_EXCEL)

                wb = openpyxl.load_workbook(NOME_FILE_EXCEL)
                if NOME_FOGLIO_TARGET in wb.sheetnames:
                    ws = wb[NOME_FOGLIO_TARGET]
                else:
                    ws = wb.active

                timestamp_attuale = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

                note_uscite_comb = f"{uscite_prob} | Altro: {uscite_altro}".strip(" | Altro:")
                note_cartello_comb = f"{cartello_lum} | Altro: {cartello_altro}".strip(" | Altro:")
                note_led_comb = f"Lampada N°: {led_lampada_num} | LED: {led_stato}" if led_lampada_num or led_stato else ""
                note_liquidi_comb = f"Litri: {liquidi_litri} | Stanze: {liquidi_stanze}" if liquidi_litri or liquidi_stanze else ""
                note_scaffali_comb = f"{scaffali_presenza} - Tipo: {scaffali_tipo}"
                note_bombole_dett = f"N° Bombole: {bombole_num} | Litri N°: {bombole_litri}" if bombole_num or bombole_litri else ""

                riga_dati = [
                    timestamp_attuale,
                    matricola,
                    dipartimento,
                    sc_ssd,
                    sede,
                    desc_sede,
                    comune,
                    via,
                    padiglione,
                    piano,
                    lato,
                    est1, est2, est3, est4, est5, est6, est7,
                    estintore_num,
                    estintore_problema,
                    note_estintori,
                    idranti,
                    idrante_num,
                    idrante_problema,
                    note_idranti,
                    coperta,
                    coperta_segnalata,
                    coperta_integra,
                    note_coperta,
                    fumi,
                    note_fumi,
                    centralina,
                    note_centralina,
                    pulsanti,
                    pulsanti_privi,
                    rei_apertura,
                    rei_maniglioni,
                    rei_num_porte,
                    rei_num_ante,
                    rei_prob,
                    non_rei_esito,
                    non_rei_num,
                    non_rei_prob,
                    non_rei_ante,
                    non_rei_stato,
                    uscite_esito,
                    note_uscite_comb,
                    note_cartello_comb,
                    planimetrie_pres,
                    planimetrie_agg,
                    lampade_integre,
                    lampade_rotte_num,
                    carica_led_pres,
                    note_led_comb,
                    liquidi_pres,
                    note_liquidi_comb,
                    armadi_integri,
                    armadi_note_libera,
                    depositi_num,
                    depositi_fumi,
                    depositi_rei,
                    depositi_estintore,
                    note_scaffali_comb,
                    bombole_pres,
                    note_bombole_dett,
                    bombole_fumi,
                    bombole_areato,
                    bombole_ancorate,
                ]

                ws.append(riga_dati)
                wb.save(NOME_FILE_EXCEL)

                # --- INVIO AUTOMATICO VIA EMAIL (INVISIBILE PER L'OPERATORE) ---
                mittente_email = st.secrets["email"]["mittente"]
                password_email = st.secrets["email"]["password"]
                destinatario_email = st.secrets["email"]["destinatario"]

                msg = EmailMessage()
                msg["Subject"] = f"Nuova Ispezione Antincendio - {dipartimento} ({sc_ssd})"
                msg["From"] = mittente_email
                msg["To"] = destinatario_email
                msg.set_content(f"È stata compilata una nuova check-list di verifica antincendio.\n\nDipartimento: {dipartimento}\nSC/SSD: {sc_ssd}\nOperatore (Matricola): {matricola}\nData e Ora: {timestamp_attuale}\n\nIn allegato trovi il file Excel aggiornato con tutte le registrazioni.")

                # Allega il file Excel aggiornato
                with open(NOME_FILE_EXCEL, "rb") as f:
                    file_data = f.read()
                    file_name = os.path.basename(NOME_FILE_EXCEL)
                msg.add_attachment(file_data, maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=file_name)

                # Invio tramite server SMTP di Google (Gmail)
                with smtplib.SMTP_SSL("smtp.gmail.com", 465) as smtp:
                    smtp.login(mittente_email, password_email)
                    smtp.send_message(msg)

                st.success("✅ Modulo compilato e inviato con successo!")

            except Exception as e:
                st.error(f"❌ Si è verificato un errore durante l'invio: {e}")